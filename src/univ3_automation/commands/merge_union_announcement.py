#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
연합광고 반영 스크립트
이번 주 행정광고 PPT에 연합광고를 병합합니다.
"""

from runtime_paths import (
    get_base_dir,
    get_dir,
    get_downloads_dir,
    get_kakao_downloads_dir,
    get_env_file_path,
    get_google_credentials_candidates,
    get_google_sheet_id,
    get_screenshot_dirs,
)

import os
import sys
from datetime import datetime, timedelta
from pptx import Presentation
import pytz
import re
import win32com.client
import pythoncom
import tempfile

# ====================================
# 설정 - 마지막에 남겨둘 슬라이드 개수
# ====================================
SLIDES_TO_KEEP_AT_END = 3

def get_this_sunday():
    """현재 한국 시간 기준으로 이번 주 일요일 날짜를 반환"""
    # 한국 시간대 설정
    kst = pytz.timezone('Asia/Seoul')
    now_kst = datetime.now(kst)

    # 현재 날짜의 요일 (0=월요일, 6=일요일)
    weekday = now_kst.weekday()

    # 이번 주 일요일 계산
    if weekday == 6:  # 이미 일요일인 경우
        sunday = now_kst
    else:
        # 다음 일요일까지 남은 일수
        days_until_sunday = 6 - weekday
        sunday = now_kst + timedelta(days=days_until_sunday)

    return sunday

def find_union_announcement_ppt(union_dir, this_sunday):
    """연합광고 디렉토리에서 이번 주 일요일 연합광고 PPT 찾기"""
    # YYMMDD 형식
    yy = this_sunday.strftime("%y")
    mmdd = this_sunday.strftime("%m%d")
    yymmdd = yy + mmdd

    best_match = None
    latest_mtime = 0

    if not os.path.exists(union_dir):
        print(f"연합광고 디렉토리를 찾을 수 없습니다: {union_dir}")
        return None

    for file in os.listdir(union_dir):
        if file.endswith('.pptx') and not file.startswith('~$'):
            # YYMMDD 연합행정광고 패턴 확인
            pattern = rf'{yymmdd}\s*연합행정광고'
            if re.search(pattern, file):
                file_path = os.path.join(union_dir, file)
                file_mtime = os.path.getmtime(file_path)

                # 가장 최근에 수정된 파일 선택
                if file_mtime > latest_mtime:
                    latest_mtime = file_mtime
                    best_match = file

    return best_match

def find_union_in_kakao(kakao_dir, this_sunday):
    """카카오톡 받은 파일에서 연합광고 PPT 찾기"""
    if not os.path.exists(kakao_dir):
        return None

    # YYMMDD 형식
    yy = this_sunday.strftime("%y")
    mmdd = this_sunday.strftime("%m%d")
    yymmdd = yy + mmdd

    best_match = None
    latest_mtime = 0

    for file in os.listdir(kakao_dir):
        if file.endswith('.pptx') and not file.startswith('~$'):
            # YYMMDD 연합행정광고 패턴 확인
            pattern = rf'{yymmdd}\s*연합행정광고'
            if re.search(pattern, file):
                file_path = os.path.join(kakao_dir, file)
                file_mtime = os.path.getmtime(file_path)

                # 가장 최근에 수정된 파일 선택
                if file_mtime > latest_mtime:
                    latest_mtime = file_mtime
                    best_match = file_path

    return best_match

def move_union_from_kakao(kakao_file_path, union_dir):
    """카카오톡 폴더에서 연합광고 폴더로 파일 이동"""
    import shutil

    # 타겟 폴더가 없으면 생성
    os.makedirs(union_dir, exist_ok=True)

    # 타겟 경로 생성
    target_path = os.path.join(union_dir, os.path.basename(kakao_file_path))

    # 파일 이동 (복사 후 삭제)
    shutil.copy2(kakao_file_path, target_path)
    os.remove(kakao_file_path)

    print(f"✓ 연합광고 파일을 이동했습니다:")
    print(f"  {os.path.basename(kakao_file_path)}")

    return os.path.basename(target_path)

def extract_ppt_to_png(ppt_path, output_dir):
    """PPT 파일의 모든 슬라이드를 PNG로 추출"""
    print(f"\nPPT를 PNG로 추출 중: {os.path.basename(ppt_path)}")

    pythoncom.CoInitialize()
    png_files = []

    try:
        powerpoint = win32com.client.Dispatch("PowerPoint.Application")
        powerpoint.Visible = 1  # PowerPoint는 반드시 Visible이어야 함
        powerpoint.WindowState = 2  # ppWindowMinimized - 최소화 상태로 실행

        # 절대 경로 사용
        abs_ppt_path = os.path.abspath(ppt_path)
        presentation = powerpoint.Presentations.Open(abs_ppt_path)

        # 슬라이드 크기 가져오기
        slide_width = presentation.PageSetup.SlideWidth
        slide_height = presentation.PageSetup.SlideHeight

        # 각 슬라이드를 PNG로 저장
        for i, slide in enumerate(presentation.Slides):
            output_file = os.path.join(output_dir, f"union_slide_{i+1:02d}.png")
            # 원본 크기로 내보내기
            slide.Export(output_file, "PNG", int(slide_width), int(slide_height))
            png_files.append(output_file)
            print(f"  - 슬라이드 {i+1} 추출 완료")

        presentation.Close()
        powerpoint.Quit()

    except Exception as e:
        print(f"PNG 추출 중 오류 발생: {str(e)}")
        import traceback
        traceback.print_exc()

    finally:
        pythoncom.CoUninitialize()

    return png_files

def merge_union_announcements(ppt_path, union_png_files):
    """부서 광고 PPT에 연합광고 이미지 병합"""
    print(f"\n연합광고 병합 시작...")

    try:
        prs = Presentation(ppt_path)

        # 현재 마지막 SLIDES_TO_KEEP_AT_END개 슬라이드 전 위치 찾기
        insert_position = len(prs.slides) - SLIDES_TO_KEEP_AT_END

        # 빈 슬라이드 레이아웃 가져오기 (빈 레이아웃 사용)
        blank_slide_layout = prs.slide_layouts[6]  # 빈 슬라이드

        # 각 PNG 파일에 대해 슬라이드 추가
        for i, png_file in enumerate(union_png_files):
            # 새 슬라이드 추가 (마지막 SLIDES_TO_KEEP_AT_END개 슬라이드 전에 삽입)
            slide = prs.slides.add_slide(blank_slide_layout)

            # 이미지 추가 (슬라이드 전체 크기로)
            left = 0
            top = 0
            pic = slide.shapes.add_picture(png_file, left, top,
                                          width=prs.slide_width,
                                          height=prs.slide_height)

            # 슬라이드를 올바른 위치로 이동
            xml_slides = prs.slides._sldIdLst
            slide_element = xml_slides[-1]  # 방금 추가한 슬라이드
            xml_slides.remove(slide_element)
            xml_slides.insert(insert_position + i, slide_element)

            print(f"  - 연합광고 슬라이드 {i+1} 추가 완료")

        # 저장
        prs.save(ppt_path)
        print(f"\n연합광고 {len(union_png_files)}개 슬라이드 병합 완료")
        return True

    except Exception as e:
        print(f"병합 중 오류 발생: {str(e)}")
        import traceback
        traceback.print_exc()
        return False

def merge_union_announcement():
    """연합광고 반영 메인 함수"""

    print("=" * 50)
    print("연합광고 반영 스크립트")
    print(f"마지막에 남겨둘 슬라이드 개수: {SLIDES_TO_KEEP_AT_END}개")
    print("=" * 50 + "\n")

    # 디렉토리 경로 설정
    parent_dir = get_base_dir()
    admin_dir = os.path.join(parent_dir, "행정광고")
    union_dir = os.path.join(admin_dir, "연합광고")

    if not os.path.exists(admin_dir):
        print(f"오류: {admin_dir} 디렉토리를 찾을 수 없습니다.")
        return False

    # 날짜 계산
    this_sunday = get_this_sunday()
    this_sunday_str = this_sunday.strftime("%m%d")  # MMDD 형식

    print(f"이번주 일요일: {this_sunday.strftime('%Y년 %m월 %d일')} ({this_sunday_str})")
    print()

    # 이번 주 행정광고 PPT 파일 경로
    ppt_name = f"{this_sunday_str} 행정광고.pptx"
    ppt_path = os.path.join(admin_dir, ppt_name)

    # PPT 파일이 존재하는지 확인
    if not os.path.exists(ppt_path):
        print(f"오류: {this_sunday_str} 행정광고.pptx 파일을 찾을 수 없습니다.")
        print("먼저 '행정광고 PPT 생성' 메뉴를 실행하세요.")
        return False

    try:
        print("연합광고 병합 프로세스를 시작합니다...\n")

        # 1. 연합광고 폴더에서 찾기
        union_ppt = find_union_announcement_ppt(union_dir, this_sunday)

        if not union_ppt:
            # 2. 카카오톡 받은 파일에서 찾기
            kakao_dir = get_kakao_downloads_dir()
            print(f"연합광고 폴더에 없습니다. 카카오톡 받은 파일에서 검색 중...")

            kakao_file = find_union_in_kakao(kakao_dir, this_sunday)

            if kakao_file:
                # 파일 이동
                union_ppt = move_union_from_kakao(kakao_file, union_dir)
            else:
                print(f"\n❌ 이번주({this_sunday.strftime('%y%m%d')}) 연합광고 PPT를 찾을 수 없습니다.")
                print(f"   연합광고 폴더: {union_dir}")
                print(f"   카카오톡 폴더: {kakao_dir}")
                print("\n연합광고 PPT 파일명 형식: YYMMDD 연합행정광고.pptx")
                print(f"예시: {this_sunday.strftime('%y%m%d')} 연합행정광고.pptx")
                return False

        print(f"✓ 연합광고 PPT 발견: {union_ppt}")

        # 임시 디렉토리에 PNG 추출
        with tempfile.TemporaryDirectory() as temp_dir:
            union_ppt_path = os.path.join(union_dir, union_ppt)
            png_files = extract_ppt_to_png(union_ppt_path, temp_dir)

            if png_files:
                # 기존 PPT에 연합광고 병합
                success = merge_union_announcements(ppt_path, png_files)

                if success:
                    print("\n연합광고 병합이 성공적으로 완료되었습니다!")
                else:
                    print("\n연합광고 병합 중 오류가 발생했습니다.")
                    return False
            else:
                print("PNG 추출에 실패했습니다.")
                return False

        print("\n작업이 성공적으로 완료되었습니다!")
        return True

    except PermissionError:
        print("\n오류: PPT 파일에 접근할 수 없습니다.")
        print("PowerPoint에서 파일이 열려있다면 닫고 다시 시도해주세요.")
        return False
    except Exception as e:
        print(f"\n오류 발생: {str(e)}")
        import traceback
        traceback.print_exc()
        return False

def main():
    """메인 실행 함수"""
    # 실행
    success = merge_union_announcement()
    sys.exit(0 if success else 1)

if __name__ == "__main__":
    main()

